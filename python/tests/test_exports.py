"""The export flavours, and the slide-unit regression they exist to prevent."""

from __future__ import annotations

import json
import re
import subprocess
from pathlib import Path

import pytest
from PIL import Image
from pptx import Presentation

from rendiv_kit.exports import SLIDE_H, SLIDE_W, ExportConfig, export
from rendiv_kit.timeline import Timeline

# Two segments, four beats. The counts differ, which is the whole point: an
# export that confuses the two units produces 2 slides where it owes 4.
TIMELINE = {
    "fps": 30,
    "width": 1920,
    "height": 1080,
    "segments": [
        {
            "id": "S1",
            "beats": [
                {"name": "records", "len": 30},
                {"name": "gaps", "len": 30},
                {"name": "claim", "len": 60},
            ],
        },
        {"id": "S2", "beats": [{"name": "close", "len": 30}]},
    ],
}


def _ffmpeg() -> str:
    from shutil import which

    found = which("ffmpeg")
    if not found:
        pytest.skip("ffmpeg is not on PATH")
    return found


@pytest.fixture
def project(tmp_path: Path) -> ExportConfig:
    root = tmp_path
    (root / "src").mkdir()
    (root / "src" / "timeline.json").write_text(json.dumps(TIMELINE), encoding="utf-8")

    stills = root / "out" / "stills"
    stills.mkdir(parents=True)
    tl = Timeline.load(root / "src" / "timeline.json")
    # Full size, so the PDF page size assertion below means something: 1920px at
    # 144 dpi is the 13.333 in the pptx slide measures.
    for beat in tl.all_beats():
        Image.new("RGB", (tl.width, tl.height), (20, 20, 20)).save(stills / beat.still_name)
    for seg in tl.segments:
        Image.new("RGB", (tl.width, tl.height), (40, 40, 40)).save(
            stills / tl.settled_still_name(seg)
        )

    return ExportConfig(root=root, stem="deck")


@pytest.fixture
def with_segments(project: ExportConfig) -> ExportConfig:
    """Segment renders, which the video flavour slices per beat."""
    ffmpeg = _ffmpeg()
    project.segments_dir.mkdir(parents=True, exist_ok=True)
    tl = Timeline.load(project.timeline)
    for seg in tl.segments:
        seconds = tl.segment_length(seg) / tl.fps
        subprocess.run(
            [
                ffmpeg, "-y", "-hide_banner", "-loglevel", "error",
                "-f", "lavfi", "-i", f"testsrc=size=192x108:rate={tl.fps}:duration={seconds}",
                "-c:v", "libx264", "-pix_fmt", "yuv420p",
                str(project.segments_dir / f"{seg['id']}.mp4"),
            ],
            check=True,
        )
    project.ffmpeg = Path(ffmpeg)
    return project


def test_pptx_video_has_one_slide_per_beat_not_per_segment(with_segments: ExportConfig):
    """The regression this flavour was rebuilt for.

    An export that reaches for a segment clip here puts all three beats of S1
    behind one advance: the room cannot be stopped between them and the
    presenter cannot step back. Four beats owe four slides.
    """
    dest = export(with_segments, "pptx-video")
    prs = Presentation(str(dest))
    assert len(prs.slides) == 4, "one slide per beat"
    assert len(prs.slides) != len(TIMELINE["segments"]), "not one slide per segment"


def test_pptx_video_embeds_a_clip_of_that_beats_length(with_segments: ExportConfig):
    """Each slide's video covers its own beat, not its segment's whole span."""
    ffmpeg = str(with_segments.ffmpeg)
    export(with_segments, "pptx-video")
    tl = Timeline.load(with_segments.timeline)

    for beat in tl.all_beats():
        clip = with_segments.out_dir / "beat-clips" / beat.clip_name
        assert clip.exists(), f"no clip cut for {beat.segment}.{beat.name}"
        probe = subprocess.run(
            [
                ffmpeg.replace("ffmpeg", "ffprobe"), "-v", "error",
                "-show_entries", "format=duration", "-of", "csv=p=0", str(clip),
            ],
            capture_output=True, text=True,
        )
        if probe.returncode != 0:
            pytest.skip("ffprobe is not available beside ffmpeg")
        expected = beat.len / tl.fps
        assert abs(float(probe.stdout.strip()) - expected) < 0.2, (
            f"{beat.segment}.{beat.name} clip is not its own beat's length"
        )


def test_pptx_video_slides_autoplay(with_segments: ExportConfig):
    """A clip left on the click trigger makes the flavour pointless."""
    dest = export(with_segments, "pptx-video")
    prs = Presentation(str(dest))
    P = "{http://schemas.openxmlformats.org/presentationml/2006/main}"
    for slide in prs.slides:
        timing = slide._element.find(P + "timing")
        assert timing is not None
        delays = [c.get("delay") for c in timing.iter(P + "cond")]
        assert "indefinite" not in delays, "media would wait for a click"


def test_pptx_stills_is_per_beat(project: ExportConfig):
    dest = export(project, "pptx-stills")
    prs = Presentation(str(dest))
    assert len(prs.slides) == 4
    assert prs.slide_width == SLIDE_W and prs.slide_height == SLIDE_H


def test_pptx_settled_is_per_segment_and_says_so_in_its_name(project: ExportConfig):
    dest = export(project, "pptx-settled")
    prs = Presentation(str(dest))
    assert len(prs.slides) == 2, "the lossy flavour, one per segment"
    assert "per-segment" in dest.name, "named for what it loses"


def test_pdf_is_per_beat_at_the_pptx_page_size(project: ExportConfig):
    dest = export(project, "pdf")
    assert dest.suffix == ".pdf"
    raw = dest.read_bytes()

    # Pillow writes PDF but cannot read one back, so the structure is read
    # directly. /Count on the page tree is the page total.
    assert re.search(rb"/Count\s+4", raw), "one page per beat"

    # 1920 / 144 dpi is 13.333 in, which is 960 pt, and 1080 / 144 is 540 pt.
    # That is exactly the pptx slide, so the two files print at one size.
    boxes = set(re.findall(rb"/MediaBox\s*\[([^\]]*)\]", raw))
    assert len(boxes) == 1, "every page the same size"
    nums = [round(float(n)) for n in boxes.pop().split()]
    assert nums == [0, 0, 960, 540], f"page is {nums[2]}x{nums[3]} pt, not the pptx slide size"


def test_beat_names_reach_the_slide_notes(project: ExportConfig):
    dest = export(project, "pptx-stills")
    prs = Presentation(str(dest))
    notes = [s.notes_slide.notes_text_frame.text for s in prs.slides]
    assert notes == ["S1 records", "S1 gaps", "S1 claim", "S2 close"]


def test_missing_stills_name_the_file_and_do_not_write_a_deck(project: ExportConfig):
    (project.stills_dir / "s1-gaps.png").unlink()
    with pytest.raises(SystemExit) as err:
        export(project, "pptx-stills")
    assert "s1-gaps" in str(err.value)


def test_missing_segment_render_names_the_segment(project: ExportConfig):
    with pytest.raises(SystemExit) as err:
        export(project, "pptx-video")
    assert "S1" in str(err.value)


def test_unknown_flavour_is_rejected(project: ExportConfig):
    with pytest.raises(SystemExit):
        export(project, "pptx-everything")
